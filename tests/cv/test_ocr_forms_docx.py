"""OCR recognition, failure handling, form extraction, evaluation metrics and DOCX generation."""
from __future__ import annotations

import json
import random

import cv2
import numpy as np
import pytest
from docx import Document
from docx.shared import Pt

from conftest import FIXTURES, gf, requires_tesseract
from export import docx_generator as dg
from ocr import form_extraction, recognition
from ocr.layout import confidence_summary, normalize_text, text_from_words
from processing import evaluation
from processing.common import ParameterError, load_image, to_gray
from processing.enhancement import enhance, resolve_params


# ------------------------------------------------------------------ OCR
@requires_tesseract
def test_ocr_reads_clean_letter_accurately(letter, reference_text):
    res = recognition.recognize(letter, psm=3)
    assert not res.empty
    assert evaluation.cer(reference_text, res.text) < 0.03
    assert res.confidence["mean"] > 85 and res.confidence["word_count"] > 80
    assert all(k in res.words[0] for k in ("text", "conf", "x", "y", "w", "h", "low"))
    assert "\n\n" in res.text  # paragraph breaks preserved


@requires_tesseract
def test_ocr_flags_low_confidence_words():
    blurred = cv2.GaussianBlur(gf.render_text_page(), (0, 0), 3.2)
    res = recognition.recognize(blurred, psm=3, threshold=95.0)
    assert res.confidence["low_confidence_count"] > 0
    assert any(w["low"] for w in res.words)


@requires_tesseract
def test_ocr_blank_image_reports_empty_with_warning():
    res = recognition.recognize(np.full((600, 800, 3), 255, np.uint8))
    assert res.empty and res.text == "" and res.warnings
    assert res.confidence["mean"] is None


def test_missing_tesseract_executable_gives_actionable_error():
    with pytest.raises(recognition.OcrUnavailableError, match="TESSERACT_CMD"):
        recognition.configure("/definitely/not/here/tesseract", force=True)
    # restore a working configuration for the following tests
    try:
        recognition.configure(None, force=True)
    except recognition.OcrUnavailableError:
        pass


@requires_tesseract
def test_ocr_timeout_is_translated(monkeypatch, letter):
    import pytesseract

    def slow(*a, **k):
        raise RuntimeError("Tesseract process timeout")

    monkeypatch.setattr(pytesseract, "image_to_string", slow)
    with pytest.raises(recognition.OcrTimeoutError, match="timed out"):
        recognition.recognize(letter, timeout=1)


@requires_tesseract
def test_ocr_engine_failure_is_translated(monkeypatch, letter):
    import pytesseract

    def fail(*a, **k):
        raise pytesseract.TesseractError(1, "boom")

    monkeypatch.setattr(pytesseract, "image_to_data", fail)
    with pytest.raises(recognition.OcrError):
        recognition.recognize(letter)


def test_layout_helpers():
    assert normalize_text("a \r\n\r\n\r\n\r\nb\x0c") == "a\n\nb"
    words = [
        {"text": "World", "x": 80, "block": 1, "par": 1, "line": 1},
        {"text": "Hello", "x": 10, "block": 1, "par": 1, "line": 1},
        {"text": "Next", "x": 10, "block": 1, "par": 2, "line": 1},
    ]
    assert text_from_words(words) == "Hello World\n\nNext"
    summary = confidence_summary([{"text": "a", "conf": 90}, {"text": "b", "conf": 40}], 60)
    assert summary["mean"] == 65 and summary["low_confidence_count"] == 1


# ------------------------------------------------------------------ DSP evaluation metrics
def _naive(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def test_edit_distance_matches_reference_dp():
    rng = random.Random(4)
    for _ in range(300):
        a = "".join(rng.choice("abc d") for _ in range(rng.randint(0, 40)))
        b = "".join(rng.choice("abc d") for _ in range(rng.randint(0, 40)))
        assert evaluation.edit_distance(a, b) == _naive(a, b)
    assert evaluation.edit_distance("kitten", "sitting") == 3


def test_cer_wer_values():
    assert evaluation.cer("hello world", "hello world") == 0
    assert evaluation.cer("abcd", "abxd") == pytest.approx(0.25)
    assert evaluation.wer("the quick brown fox", "the quick brown dog") == pytest.approx(0.25)
    assert evaluation.wer("a b", "a") == pytest.approx(0.5)
    assert evaluation.cer("", "anything") is None  # no ground truth -> no accuracy number
    assert evaluation.cer("a  b\n c", "a b c") == 0  # whitespace is normalised


@requires_tesseract
def test_experiments_show_processing_helps_on_a_photograph(letter_photo, reference_text):
    out = evaluation.run_experiments(letter_photo, ["A", "D"], enhance_preset="document_scan", reference_text=reference_text)
    rows = {r["config"]: r for r in out["results"]}
    assert rows["A"]["cer"] is not None and rows["D"]["cer"] is not None
    assert rows["D"]["cer"] < rows["A"]["cer"]
    assert rows["D"]["cer"] < 0.08
    assert rows["D"]["geometry_ms"] > 0 and rows["D"]["enhancement_ms"] > 0 and rows["A"]["geometry_ms"] == 0


@requires_tesseract
def test_experiments_without_reference_do_not_invent_accuracy(letter):
    out = evaluation.run_experiments(letter, ["A"], reference_text=None)
    assert out["results"][0]["cer"] is None and out["results"][0]["wer"] is None and not out["reference_provided"]


# ------------------------------------------------------------------ form extraction
def test_template_is_well_formed():
    tpl = form_extraction.load_template("university_registration_v1")
    keys = [f["key"] for f in tpl["fields"]]
    assert keys == ["student_name", "student_number", "program", "year_level", "section", "date", "contact_number", "address"]
    for f in tpl["fields"]:
        x, y, w, h = f["region"]
        assert 0 <= x and 0 <= y and x + w <= 1 and y + h <= 1
    with pytest.raises(ParameterError):
        form_extraction.load_template("../../etc/passwd")
    with pytest.raises(ParameterError):
        form_extraction.load_template("nope_v9")


@requires_tesseract
def test_form_fields_are_extracted_from_clean_form(form_clean):
    result = form_extraction.extract_fields(form_clean, "university_registration_v1")
    expected = json.loads((FIXTURES / "registration_form_reference.json").read_text())
    assert result["matched"] and result["match_score"] == 1.0
    got = {f["key"]: f["value"] for f in result["fields"]}
    assert got == expected


@requires_tesseract
def test_form_extraction_from_photographed_form(form_photo):
    from processing.pipeline import run_geometry

    corrected, _ = run_geometry(form_photo)
    params, _ = resolve_params("document_scan", None, to_gray(corrected))
    enhanced, _ = enhance(corrected, params)
    result = form_extraction.extract_fields(enhanced, "university_registration_v1")
    expected = json.loads((FIXTURES / "registration_form_reference.json").read_text())
    assert result["matched"]
    got = {f["key"]: f["value"] for f in result["fields"]}
    correct = sum(1 for k, v in expected.items() if got.get(k) == v)
    assert correct >= 7, f"only {correct}/8 fields exact: {got}"


@requires_tesseract
def test_unknown_document_does_not_match_template(letter):
    result = form_extraction.extract_fields(letter, "university_registration_v1")
    assert not result["matched"] and result["fields"] == [] and "does not look like" in result["message"]


# ------------------------------------------------------------------ DOCX
def _all_runs(doc):
    runs = [r for p in doc.paragraphs for r in p.runs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    runs.extend(p.runs)
    return runs


def test_text_docx_formatting_and_content(tmp_path):
    path = tmp_path / "t.docx"
    text = "First paragraph wraps\nonto a second line.\n\nSecond paragraph.\n\n\nThird: ünïcode – “quotes” ✓"
    info = dg.build_text_document(path, text, reflow=True, title="My Title")
    doc = Document(str(path))
    runs = _all_runs(doc)
    assert runs and all(r.font.name == "Times New Roman" and r.font.size == Pt(12) for r in runs)
    assert doc.styles["Normal"].font.name == "Times New Roman" and doc.styles["Normal"].font.size == Pt(12)
    sec = doc.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0 and round(sec.page_height.cm, 1) == 29.7
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs[0] == "My Title"
    assert "First paragraph wraps onto a second line." in paragraphs  # reflowed
    assert "Third: ünïcode – “quotes” ✓" in paragraphs  # no encoding damage
    assert not doc.inline_shapes  # text is text, not an embedded picture
    assert info["paragraphs"] == 3


def test_text_docx_without_reflow_keeps_line_breaks(tmp_path):
    path = tmp_path / "t.docx"
    dg.build_text_document(path, "line one\nline two", reflow=False)
    doc = Document(str(path))
    assert doc.paragraphs[0].text == "line one\nline two"


def test_docx_strips_illegal_control_characters(tmp_path):
    path = tmp_path / "t.docx"
    dg.build_text_document(path, "ok\x00 text\x0b here\x0c!")
    assert Document(str(path)).paragraphs[0].text.replace("\n", "") == "ok text here!"


def test_empty_text_still_creates_valid_document(tmp_path):
    path = tmp_path / "t.docx"
    dg.build_text_document(path, "")
    assert dg.verify_docx(path)["font_ok"]


def test_form_docx_contains_editable_table(tmp_path):
    path = tmp_path / "f.docx"
    fields = [{"key": "student_name", "label": "Student Name", "value": "DELA CRUZ, JUAN M."}, {"key": "address", "label": "Address", "value": "Line 1\nLine 2"}]
    info = dg.build_form_document(path, "Student Registration Form", fields, full_text="Other text here.", include_full_text=True)
    doc = Document(str(path))
    assert len(doc.tables) == 1
    rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert rows[0] == ["Field", "Value"] and rows[1] == ["Student Name", "DELA CRUZ, JUAN M."]
    assert rows[2][1] == "Line 1\nLine 2"
    assert doc.paragraphs[0].text == "STUDENT REGISTRATION FORM"
    assert any("Other text here." in p.text for p in doc.paragraphs)
    assert all(r.font.name == "Times New Roman" and r.font.size == Pt(12) for r in _all_runs(doc))
    assert info["table_rows"] == 3
    with pytest.raises(ParameterError):
        dg.build_form_document(tmp_path / "x.docx", "T", [])


def test_docx_defaults_do_not_keep_theme_fonts(tmp_path):
    path = tmp_path / "t.docx"
    dg.build_text_document(path, "hello")
    import zipfile

    styles = zipfile.ZipFile(path).read("word/styles.xml").decode("utf8")
    assert "asciiTheme" not in styles.split("<w:style ")[0]  # docDefaults section
    assert 'w:ascii="Times New Roman"' in styles


@requires_tesseract
@pytest.mark.parametrize("sigma,scale", [(16, 1), (10, 2), (14, 2)])
def test_automatic_preset_keeps_noisy_pages_readable(letter, reference_text, sigma, scale):
    """Regression: CLAHE at its default strength used to amplify noise until Tesseract returned nothing at all."""
    gray = cv2.resize(to_gray(letter), (880, 1245), interpolation=cv2.INTER_AREA)
    rng = np.random.default_rng(11)
    h, w = gray.shape
    field = rng.normal(0, sigma * scale, (h // scale + 1, w // scale + 1)).astype(np.float32)
    field = cv2.resize(field, (w, h), interpolation=cv2.INTER_LINEAR) if scale > 1 else rng.normal(0, sigma, gray.shape).astype(np.float32)
    noisy = np.clip(gray.astype(np.float32) + field, 0, 255).astype(np.uint8)
    params, _ = resolve_params("automatic", None, noisy)
    out, _ = enhance(cv2.cvtColor(noisy, cv2.COLOR_GRAY2BGR), params)
    res = recognition.recognize(out, psm=3)
    assert evaluation.cer(reference_text, res.text) < 0.15


@pytest.mark.parametrize(
    "value,ftype,plausible",
    [
        ("DELA CRUZ, JUAN M.", "name", True), ("O'Brien-Smith, Mary Ann", "name", True), ("José Núñez", "name", True),
        ("DELA CRUZ, |UAN M.", "name", False), ("Juan 2nd", "name", False), ("", "name", False),
        ("2022-01457", "id", True), ("20", "id", False),
        ("0917-123-4567", "phone", True), ("123", "phone", False),
        ("15/09/2026", "date", True), ("15 September 2026", "date", True), ("September 15, 2026", "date", True), ("2026-09-15", "date", True), ("foo", "date", False),
        ("BS COMPUTER ENGINEERING", "text", True), ("CPE-4A", "text", True), ("BL| 5 LOT", "text", False),
        ("BLK 5 LOT 12 SAMPAGUITA ST. SAN PEDRO CITY LAGUNA 4023", "multiline", True),
    ],
)
def test_field_content_checks_flag_misread_characters(value, ftype, plausible):
    """A letter read as a symbol (J -> '|') can still come with high engine confidence, so the content is checked too."""
    assert (form_extraction.check_plausible(value, ftype) is None) is plausible
